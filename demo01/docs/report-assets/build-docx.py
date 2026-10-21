#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""根据空白模板格式生成银行存取款管理系统的实验报告 docx。"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parents[2]      # demo01
ASSETS = BASE / "docs" / "report-assets"
OUT = BASE / "docs" / "实验报告-银行存取款管理系统.docx"

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
CODE_FONT = "Consolas"


def set_run_font(run, name=BODY_FONT, size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def body(text, indent=True, size=10.5, space_after=3, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.4
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    run = p.add_run(text)
    set_run_font(run, BODY_FONT, size, bold)
    return p


def bullet(text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_run_font(run, BODY_FONT, size)
    return p


def numbered(text, size=10.5):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_run_font(run, BODY_FONT, size)
    return p


def heading(text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    size = {1: 15, 2: 12.5, 3: 11}[level]
    run = p.add_run(text)
    set_run_font(run, HEAD_FONT, size, True, (0x1F, 0x3A, 0x5F))
    return p


def shade(paragraph, fill="ECEFF3"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "4")
    left.set(qn("w:color"), "9FB6CD")
    pBdr.append(left)
    pPr.append(pBdr)


def code(text, size=8):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    shade(p)
    run = p.add_run(text.rstrip("\n"))
    set_run_font(run, CODE_FONT, size)
    return p


def figure(filename, caption, width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ASSETS / filename), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    run = cap.add_run(caption)
    set_run_font(run, BODY_FONT, 9, False, (0x2D, 0x3D, 0x4F))


def cell_shade(cell, fill="F2F4F7"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, size=10, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, BODY_FONT, size, bold)


def table_borders(table, color="9AA7B4", size=6):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def table_margins(table, left=100, right=100, top=40, bottom=40):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def info_table(rows):
    table = doc.add_table(rows=len(rows), cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(table)
    table_margins(table)
    widths = [Cm(2.6), Cm(5.4), Cm(2.6), Cm(5.4)]
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
    for r, row in enumerate(rows):
        if len(row) == 2:      # 跨三列的长值
            cell_text(table.cell(r, 0), row[0], True, 10, True)
            cell_shade(table.cell(r, 0))
            merged = table.cell(r, 1).merge(table.cell(r, 3))
            cell_text(merged, row[1], False, 10)
        else:
            for c, (text, is_label) in enumerate(row):
                cell_text(table.cell(r, c), text, is_label, 10, is_label)
                if is_label:
                    cell_shade(table.cell(r, c))
    table.rows[5].height = Cm(2.2)
    table.rows[6].height = Cm(1.2)
    return table


doc = Document()

# 页面与正文样式：A4，适合打印
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)

normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

# ---------- 封面（按空白模板字段） ----------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(12)
title.paragraph_format.space_after = Pt(4)
set_run_font(title.add_run("实验报告"), HEAD_FONT, 22, True)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(18)
set_run_font(sub.add_run("学院（系）名称：华信软件学院"), BODY_FONT, 11)

info_table([
    [("姓名", True), ("张欣悦", False), ("学号", True), ("20241779", False)],
    [("专业", True), ("软件工程", False), ("班级", True), ("2班", False)],
    ("实验项目", "实验一 MyBatis Lab（银行存取款管理系统）"),
    [("课程名称", True), ("JAVA 框架技术（一）", False), ("课程代码", True), ("1960229", False)],
    [("实验时间", True), ("2026.9—2026.10", False), ("实验地点", True), ("机房", False)],
    [("批改意见", True), ("", False), ("成绩", True), ("", False)],
    ("教师签字", "　　　　　　　　　　　　　　　　　　　　　日期："),
])

doc.add_page_break()

# ---------- 一、实验概况 ----------
heading("一、实验概况")
body("实验名称：实验一 MyBatis Lab（自拟业务：银行存取款管理系统）。实验时间 2026 年 9 月—10 月，共 4 学时。")
body("实验环境：macOS；JDK 21（OpenJDK）；VS Code + GitHub Copilot；Maven 3.9.6（阿里云镜像与本地仓库）；MySQL 8.4.9；MyBatis 3.5.16、MyBatis-Plus 3.5.7、MySQL Connector/J 8.4.0、Log4j 1.2.17、JUnit Jupiter 5.10.2。")
body("版本说明：实验指导书中的 ssm_emp、emp 表及 JDK 8/11、JUnit 4 为参考内容；经老师确认，个人项目可自拟业务与版本。本项目采用银行三表模型与现代 JDK/JUnit 5。")
body("业务背景：数据库 mybatis_demo，包含客户表 bank_customer、账户表 bank_account、流水表 bank_transaction。账户通过 customer_id 外键关联客户，流水通过 account_id 外键关联账户；余额与流水均有 CHECK 约束。")
body("功能范围：控制台菜单共 10 项——查询客户、新增客户、修改客户、开户、余额查询、存款、取款、流水分页、客户账单、退出系统。存款/取款在同一事务内完成行锁、余额更新与流水写入。")

# ---------- 二、实验目标 ----------
heading("二、实验目标")
for text in [
    "掌握基于 Maven 搭建 MyBatis 环境的流程，理解全局配置、类型别名、数据源与 Mapper 注册的规范，能排查依赖冲突。",
    "理解 SqlSessionFactory、SqlSession 与 Mapper 代理的关系，掌握用 MyBatisUtil 维护全局工厂、按业务操作管理会话。",
    "熟练使用 Mapper 接口与 XML 实现 CRUD，掌握 #{} 参数绑定、自定义 resultMap、<sql>/<include> 列复用与自增主键回填。",
    "掌握 if、trim、set、foreach、choose/when/otherwise 等动态 SQL 标签，理解 @Param 命名与批量操作。",
    "掌握 MyBatis-Plus 的 BaseMapper、Lambda 条件构造器、分页插件与物理分页日志。",
    "掌握 Log4j 日志追踪与 JUnit 单元测试，能依据日志定位参数绑定、结果映射与主键回填问题。",
]:
    numbered(text)

# ---------- 三、实验内容与过程 ----------
heading("三、实验内容与过程")

heading("任务 1　数据库与工程初始化", 2)
body("三张表通过外键形成“客户 → 账户 → 流水”的层级关系，金额统一使用 DECIMAL(18,2)，并用 CHECK 约束保证余额不为负、流水金额与前后余额一致。")
code("""CREATE TABLE bank_customer (
  id INT PRIMARY KEY AUTO_INCREMENT,
  name VARCHAR(50) NOT NULL,
  phone VARCHAR(30) NOT NULL
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;

CREATE TABLE bank_account (
  id BIGINT PRIMARY KEY AUTO_INCREMENT,
  account_no VARCHAR(34) NOT NULL,
  customer_id INT NOT NULL,
  balance DECIMAL(18,2) NOT NULL DEFAULT 0.00,
  created_at DATETIME(6) NOT NULL DEFAULT CURRENT_TIMESTAMP(6),
  UNIQUE KEY uk_bank_account_no (account_no),
  CONSTRAINT fk_bank_account_customer FOREIGN KEY (customer_id)
    REFERENCES bank_customer(id) ON DELETE RESTRICT ON UPDATE RESTRICT,
  CONSTRAINT chk_bank_account_balance CHECK (balance >= 0)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;

CREATE TABLE bank_transaction (
  id BIGINT PRIMARY KEY AUTO_INCREMENT,
  account_id BIGINT NOT NULL,
  type VARCHAR(8) NOT NULL,
  amount DECIMAL(18,2) NOT NULL,
  balance_before DECIMAL(18,2) NOT NULL,
  balance_after DECIMAL(18,2) NOT NULL,
  created_at DATETIME(6) NOT NULL DEFAULT CURRENT_TIMESTAMP(6),
  CONSTRAINT fk_bank_transaction_account FOREIGN KEY (account_id)
    REFERENCES bank_account(id) ON DELETE RESTRICT ON UPDATE RESTRICT,
  CONSTRAINT chk_bank_transaction CHECK (
    amount > 0 AND balance_before >= 0 AND balance_after >= 0 AND
    ((type='DEPOSIT' AND balance_after=balance_before+amount) OR
     (type='WITHDRAW' AND balance_after=balance_before-amount)))
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;""")

body("Maven 依赖：MyBatis-Plus 3.5.7 内部依赖 MyBatis 核心；非 Spring 环境下部分默认方法需要 spring-core 的泛型解析能力。通过 dependency:tree 确认没有重复的 MyBatis 版本冲突。")
code("""<dependency>org.mybatis:mybatis:3.5.16</dependency>
<dependency>com.baomidou:mybatis-plus:3.5.7</dependency>
<dependency>com.mysql:mysql-connector-j:8.4.0</dependency>
<dependency>log4j:log4j:1.2.17</dependency>
<dependency>org.junit.jupiter:junit-jupiter:5.10.2 (test)</dependency>""")

body("工程结构：entity / mapper / service / util / org.example 分包；resources 下存放 mybatis-config.xml、db.properties、log4j.properties 与 mapper XML；test 下 6 个测试类。项目路径下共 8 次 Git 提交（2026-09-17 至 2026-10-22），每个任务节点完成后提交。")
figure("fig-git.png", "图 1　Git 提交记录（8 次提交）")

heading("任务 2　基础 CRUD、Log4j 与 JUnit", 2)
body("全局配置开启驼峰映射与 LOG4J 日志实现，为实体包设置短别名；数据源参数抽取到 db.properties，通过 ${} 占位引用。")
code("""<settings>
  <setting name="vfsImpl" value="com.example.util.VisibleFilesVfs"/>
  <setting name="mapUnderscoreToCamelCase" value="true"/>
  <setting name="lazyLoadingEnabled" value="true"/>
  <setting name="aggressiveLazyLoading" value="false"/>
  <setting name="logImpl" value="LOG4J"/>
</settings>
<typeAliases><package name="com.example.entity"/></typeAliases>
<environments default="development">
  <environment id="development">
    <transactionManager type="JDBC"/>
    <dataSource type="POOLED">
      <property name="driver" value="${jdbc.driver}"/>
      <property name="url" value="${jdbc.url}"/>
      <property name="username" value="${jdbc.username}"/>
      <property name="password" value="${jdbc.password}"/>
    </dataSource>
  </environment>
</environments>""")

body("CustomerMapper 同时继承 BaseMapper（通用 CRUD）并保留自定义 XML 方法；三张表均定义自定义 resultMap，查询列通过 <sql id=\"allColumns\"> + <include> 复用。")
code("""public interface CustomerMapper extends BaseMapper<Customer> {
  List<Customer> selectByCondition(@Param("condition") Customer condition);
  List<Customer> selectByPriority(@Param("name") String name, @Param("phone") String phone);
  int updateFields(@Param("customer") Customer customer);
  int insertRows(@Param("customers") List<Customer> customers);
  int deleteRows(@Param("ids") List<Integer> ids);
}

<sql id="allColumns">id, name, phone</sql>
<resultMap id="customerMap" type="Customer">
  <id property="id" column="id"/>
  <result property="name" column="name"/>
  <result property="phone" column="phone"/>
</resultMap>
<insert id="insertRows" useGeneratedKeys="true" keyProperty="customers.id">
  INSERT INTO bank_customer (name, phone) VALUES
  <foreach collection="customers" item="customer" separator=",">
    (#{customer.name}, #{customer.phone})
  </foreach>
</insert>""")

body("MyBatisUtil 维护全局唯一的 SqlSessionFactory；分页拦截器在工厂构建后以代码注册，避免写进 XML <plugins> 时无法注入内部拦截器列表。")
code("""private static final SqlSessionFactory FACTORY = createSqlSessionFactory(new Properties());

public static SqlSessionFactory createSqlSessionFactory(Properties properties) {
  try (InputStream in = Resources.getResourceAsStream("mybatis-config.xml")) {
    SqlSessionFactory factory = new MybatisSqlSessionFactoryBuilder().build(in, properties);
    MybatisPlusInterceptor interceptor = new MybatisPlusInterceptor();
    interceptor.addInnerInterceptor(new PaginationInnerInterceptor(DbType.MYSQL));
    factory.getConfiguration().addInterceptor(interceptor);
    return factory;
  } catch (IOException e) { throw new RuntimeException("加载配置失败", e); }
}""")

body("控制台运行结果：菜单、条件查询 SQL 与结果框。日志中的 Preparing 显示带 ? 的 SQL，Parameters 显示绑定实参，Total 显示返回行数。")
figure("fig-console.png", "图 2　控制台菜单与条件查询运行结果")

body("JUnit 测试共 6 个测试类、42 个用例。每个测试类创建独立的 bank_test_ 数据库，结束后删除；写操作用例通过回滚或临时库隔离，保证可重复执行。")
figure("fig-test.png", "图 3　mvn test：42 个用例全部通过")

heading("任务 3　动态 SQL 与批量操作", 2)
body("用 <trim prefix=\"WHERE\" prefixOverrides=\"AND |OR \"> 等效实现 <where>：有条件时自动补 WHERE 并去掉开头多余的 AND，条件全空时不生成 WHERE。")
code("""<select id="selectByCondition" resultMap="customerMap">
  SELECT <include refid="allColumns"/> FROM bank_customer
  <trim prefix="WHERE" prefixOverrides="AND |OR ">
    <if test="condition != null and condition.name != null and !condition.name.isBlank()">
      AND name LIKE CONCAT('%', #{condition.name}, '%')
    </if>
    <if test="condition != null and condition.phone != null and !condition.phone.isBlank()">
      AND phone LIKE CONCAT('%', #{condition.phone}, '%')
    </if>
  </trim>
  ORDER BY id
</select>""")

body("分支查询用 choose/when/otherwise 实现“姓名精确 → 电话 → 全部”；选择性更新用 set 去掉末尾逗号；批量插入/删除用 foreach。")
code("""<choose>
  <when test="name != null and !name.isBlank()">WHERE name = #{name}</when>
  <when test="phone != null and !phone.isBlank()">WHERE phone = #{phone}</when>
  <otherwise><!-- 查询全部 --></otherwise>
</choose>

<update id="updateFields">
  UPDATE bank_customer
  <set>
    <if test="customer.name != null">name = #{customer.name},</if>
    <if test="customer.phone != null">phone = #{customer.phone},</if>
  </set>
  WHERE id = #{customer.id}
</update>

<delete id="deleteRows">
  DELETE FROM bank_customer WHERE id IN
  <foreach collection="ids" item="id" open="(" separator="," close=")">#{id}</foreach>
</delete>""")

heading("任务 4　MyBatis-Plus 与分页", 2)
body("CustomerMapper 继承 BaseMapper 后获得 insert/selectById/updateById/deleteById/selectPage 等通用方法；客户分页使用 Lambda 条件构造器，方法引用避免列名拼写错误。")
code("""default Page<Customer> selectCustomerPage(long current, long size, String name, String phone) {
  BankValidation.page(current, size);
  return selectPage(new Page<>(current, size), Wrappers.<Customer>lambdaQuery()
      .like(name != null && !name.isBlank(), Customer::getName, name)
      .like(phone != null && !phone.isBlank(), Customer::getPhone, phone)
      .orderByAsc(Customer::getId));
}""")

body("分页日志先出现 COUNT(*) AS total，再出现带 LIMIT ? 的查询，参数 10(Long) 为每页条数，说明分页拦截器真实生效且为物理分页。")
figure("fig-page.png", "图 4　分页 SQL：COUNT + LIMIT")

heading("任务 5　三表关联、事务一致性与并发（拓展）", 2)
body("“客户账单”按客户编号一次查询全部账户；“流水分页”先按账号定位账户，再按 account_id 分页查询流水，避免逐行查询造成的 N+1。存款/取款在同一事务中完成行锁、余额更新与流水写入。")
code("""return transaction(s -> {
  BankAccountMapper mapper = s.getMapper(BankAccountMapper.class);
  BankAccount account = requireAccount(mapper.selectForUpdate(number)); // 行锁
  BigDecimal before = account.getBalance();
  BigDecimal after = type == TransactionType.DEPOSIT
      ? before.add(money) : before.subtract(money);
  if (after.signum() < 0) throw new IllegalArgumentException("余额不足");
  if (after.compareTo(BankValidation.MAX_BALANCE) > 0)
    throw new IllegalArgumentException("存款后余额超出允许范围");
  if (mapper.updateBalance(account.getId(), after) != 1)
    throw new IllegalStateException("余额更新失败");
  BankTransaction entry = new BankTransaction();
  entry.setAccountId(account.getId());
  entry.setType(type);
  entry.setAmount(money);
  entry.setBalanceBefore(before);
  entry.setBalanceAfter(after);
  if (s.getMapper(BankTransactionMapper.class).insert(entry) != 1)
    throw new IllegalStateException("流水写入失败");
  return s.getMapper(BankTransactionMapper.class).selectById(entry.getId());
});""")
body("查询使用 SELECT ... FOR UPDATE 锁定账户行；余额不足、溢出或写入失败都会回滚，失败操作不产生流水。并发测试用 3 个用例验证多线程扣款不会超扣。")

# ---------- 四、常见问题与排查 ----------
heading("四、常见问题与排查")

def grid(headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(table)
    table_margins(table)
    for c, h in enumerate(headers):
        cell_text(table.cell(0, c), h, True, 9.5, True)
        cell_shade(table.cell(0, c))
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            cell_text(table.cell(r, c), text, False, 9)
    return table

grid(["错误现象", "可能原因与处理"], [
    ["BindingException: Parameter 'xxx' not found",
     "XML 占位符与接口 @Param 名称不一致；查看异常中的 Available parameters，统一命名。"],
    ["查询有行数但字段为 null",
     "resultMap 的 column 与查询列名不一致，或驼峰映射未开启；对照列清单检查。"],
    ["插入成功但主键为 null",
     "缺少 useGeneratedKeys=\"true\" 或 keyProperty 路径错误；批量插入指向集合元素属性。"],
    ["分页返回全量数据",
     "分页拦截器未注册；本项目在 MyBatisUtil 中注册 MybatisPlusInterceptor + PaginationInnerInterceptor。"],
    ["log4j.properties 中文前缀乱码",
     "Log4j 1.x 按 ISO-8859-1 读取 properties；中文改用 Unicode 转义，如 [\\u65E5\\u5FD7]。"],
    ["删除客户时抛外键约束异常",
     "客户名下仍有账户，fk_bank_account_customer 为 ON DELETE RESTRICT；先处理账户或保留客户。"],
    ["macOS 外部卷 ._* 文件干扰扫描",
     "AppleDouble 元数据被 VFS 扫描；项目使用 VisibleFilesVfs，Surefire 排除 **/._*。"],
])

# ---------- 五、错误演练与 ApiFox ----------
heading("五、错误演练与 ApiFox")

heading("1. 基于日志的三类错误演练", 2)
body("演练使用临时数据库 bank_fault_drill 与临时 classpath 覆盖故障版 Mapper，不修改 src 源码；每次演练后删除临时库并重跑测试。完整日志与复现脚本见 docs/fault-drill/。")

body("（1）参数名写错：把 #{accountNo} 改成 #{accountno} 后，SQL 尚未发送就抛异常，日志直接给出可用参数列表。")
code("""org.apache.ibatis.binding.BindingException: Parameter 'accountno' not found.
Available parameters are [accountNo, param1]""")
figure("fig-err1.png", "图 5　参数名写错的运行日志")
body("排查结论：异常第二行列出可用参数，据此把 XML 占位符改回 #{accountNo}；参数绑定大小写敏感。")

body("（2）resultMap 列名写错：把 column=\"account_no\" 改成 column=\"account_no_typo\"，SQL 正常执行且 Total: 1，但账号静默变为 null。")
code("""[日志] BankAccountMapper.selectByAccountNo - <==      Total: 1
查询结果：BankAccount{id=1, accountNo=null, customerId=1, balance=0.00, ...}
accountNo = null""")
figure("fig-err2.png", "图 6　resultMap 列名写错的运行日志")
body("排查结论：映射错误不抛异常，只能通过“有行数但字段为 null”发现；需对照 resultMap 的 column 与查询列名。")

body("（3）去掉 useGeneratedKeys：插入成功（Updates: 1），但回填主键为 null。")
code("""[日志] CustomerMapper.insertRows - <==    Updates: 1
影响行数：1
回填主键：null""")
figure("fig-err3.png", "图 7　主键未回填的运行日志")
body("排查结论：自增主键回填依赖 useGeneratedKeys=\"true\" 与 keyProperty；批量插入的 keyProperty 指向集合元素属性（customers.id）。")

heading("2. ApiFox 与 AI 代码甄别", 2)
body("ApiFox 基础练习本次未执行，报告不虚构请求与响应结果，该项待后续补充截图。")
body("AI 代码审查重点：依赖版本兼容性、SQL 列名与表结构一致性、resultMap 是否覆盖全部字段、动态标签的空值边界、是否误用 ${} 拼接用户输入。本项目 SQL 映射未使用 ${}，用户输入全部通过 #{} 绑定。")

# ---------- 六、AI 应用记录 ----------
heading("六、AI 应用记录")
for text in [
    "项目初始化与持久层实现：AI 辅助生成 Maven 依赖清单、MyBatis/MP 配置结构、实体与 Mapper/XML 骨架；人工运行 dependency:tree 与测试，确认无 MyBatis 重复依赖、SQL 与表结构一致。README 注明实现与测试由 OpenAI Codex 协助。",
    "配置补齐与输出可读性：AI 辅助补齐 typeAliases、三张表的 resultMap、<sql>/<include> 与 <trim>，并调整控制台菜单顺序与结果框线；每次改动后执行 42 个用例。",
    "故障演练：AI 辅助设计“临时库 + 临时 classpath 覆盖故障 XML”的演练方式；人工实际运行三组演练并保存日志与截图，确认三类现象。",
    "质量评价：AI 能快速给出配置骨架与排查方向，但对非 Spring 环境、插件注册方式等约束把握不稳；所有输出以 mvn test、SQL 日志与数据库约束为最终判据。",
]:
    bullet(text)

# ---------- 七、思考题 ----------
heading("七、思考题")

heading("1. #{} 和 ${} 的区别", 2)
body("#{} 使用 PreparedStatement 预编译，SQL 中生成 ? 占位符，参数单独绑定，可防止 SQL 注入；${} 做字符串拼接，用户输入会直接进入 SQL。普通参数一律用 #{}；只有表名、列名、排序字段等无法参数化的位置才考虑 ${}，且必须做白名单校验。本项目所有 SQL 映射只使用 #{}。")

heading("2. SqlSessionFactory、SqlSession 和 Mapper 的关系及生命周期", 2)
body("Factory 由配置构建，创建成本高，应用级共享一个；Session 表示一次数据库会话，线程不安全，用完必须关闭，本项目在 Service 的 transaction() 中 try-with-resources 管理；Mapper 由 session.getMapper() 获取，是动态代理对象，方法调用被转换成 SQL statement。作用域依次为应用级、操作级、方法级。")

heading("3. 什么是 N+1 查询问题", 2)
body("先查 1 次得到 N 条主记录，再对每条记录各发 1 次关联查询，总共 1+N 次。本项目若先查客户列表、再逐个客户查账户就是 N+1；实际实现用 selectByCustomerId 一次取回全部账户，流水按 account_id 一次分页查询，避免逐行查询。也可用 JOIN 或内存分组。延迟加载只是推迟查询时机，不减少次数。")

heading("4. XML 动态 SQL 和 MyBatis-Plus 条件构造器的取舍", 2)
body("XML 适合复杂条件、批量语句、多表 JOIN 与精确控制 SQL；本项目客户多条件查询、批量插入/删除、流水分页均用 XML。MP 条件构造器写法简洁，Lambda 方法引用有编译期检查，适合单表 CRUD 与普通条件分页。取舍原则：单表清晰查询用 MP，复杂动态条件或批量语句用 XML。")

heading("5. AI 生成代码可能出现的问题及人工校验", 2)
body("常见问题包括：配置形式正确但运行无效（如分页拦截器写进 XML）、忽略运行环境（非 Spring 下的 MP 配置）、resultMap 列名不一致导致静默 null、空值与边界处理不足、并发修改同一文件造成语法损坏。校验手段：实际启动、mvn test、SQL 日志、断言与数据库约束。")

# ---------- 八、总结 ----------
heading("八、实验总结与思考")
body("本次实验把配置加载、工厂构建、会话管理、Mapper 代理执行 SQL 的完整链路串了起来。基础 CRUD 需要同时对齐实体属性、接口方法与 XML statement 三处命名；动态 SQL 要逐条考虑空值边界，trim、set、foreach 分别解决多余 AND、末尾逗号与批量拼值。接入 MyBatis-Plus 后，BaseMapper 与 Lambda 条件构造器减少了模板代码，但前提是工厂使用 MybatisSqlSessionFactoryBuilder 并正确注册分页拦截器。")
body("银行三表模型让我体会到数据库约束与事务边界的重要性：账户行锁、余额更新与流水写入必须处于同一事务，否则会出现余额与流水不一致；CHECK 约束与外键在数据库层兜底。调试中也确认了“编译通过不等于功能正确”：resultMap 列名错误与主键未回填都不报错，只能靠日志、断言与数据核对发现。")
body("后续可补充批量插入边界、并发锁等待、分页越界与排序白名单用例；ApiFox 请求练习尚未完成，可在后续实验中把控制台入口改造成 REST 接口后完成联调。")

# 页脚
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(footer.add_run("天津理工大学教务处制表"), BODY_FONT, 9, False, (0x55, 0x55, 0x55))

doc.save(OUT)
print("saved:", OUT)
print("paragraphs:", len(doc.paragraphs), "images:", len(doc.inline_shapes))
